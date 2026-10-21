# -*- coding: utf-8 -*-
"""
search_core.py — Excel / CSV 關鍵字搜尋核心

不依賴 PyQt,可被呼叫,也可直接用命令列執行:
    python search_core.py "C:\\Users\\User\\Desktop" result.xlsx -k "google chrome" -k "microsoft edge"

需要套件:  pip install pandas numpy openpyxl xlrd
"""
from __future__ import annotations

import multiprocessing
import os
import re
import time
import unicodedata
from concurrent.futures import FIRST_COMPLETED, ProcessPoolExecutor, wait
from concurrent.futures.process import BrokenProcessPool
from dataclasses import dataclass
from typing import Callable, Iterator, Optional

import numpy as np
import pandas as pd

ALL_EXTS = (".csv", ".xlsx", ".xls")
EXCEL_MAX_ROWS = 1_048_575          # 扣掉標題列
CSV_CHUNK_ROWS = 200_000            # 大型 CSV 分塊讀取,避免一次吃光記憶體
POOL_MIN_FILES = 6                  # 檔案數少於此值就不開多核心 (啟動成本不划算)
META_COLUMNS = ("檔案路徑", "檔案名稱", "來源工作表", "資料列序", "命中關鍵字")
_SEP = "\x1f"                       # 串接整列時的分隔字元,避免關鍵字跨儲存格誤判
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# ───────────────────────────── 資料結構 ─────────────────────────────
class SearchError(Exception):
    """可直接顯示給使用者看的錯誤。"""


@dataclass(frozen=True)
class SearchOptions:
    keywords: tuple
    case_sensitive: bool = False
    whole: bool = False
    whole_mode: str = "cell"        # "cell": 整個儲存格相同 / "word": 完整單字
    match_all: bool = False         # True: 全部關鍵字都要出現;False: 任一個即可
    normalize_width: bool = False   # True: 全形半形視為相同
    exts: tuple = ALL_EXTS
    recursive: bool = True
    skip_hidden: bool = True


@dataclass
class FileResult:
    path: str
    frame: Optional[pd.DataFrame] = None
    error: Optional[str] = None


@dataclass
class SearchSummary:
    cancelled: bool
    scanned: int
    matched_files: int
    rows: int
    elapsed: float
    output: Optional[str] = None
    preview: Optional[pd.DataFrame] = None


# ───────────────────────────── 檔案蒐集 ─────────────────────────────
def _is_hidden(full: str, name: str) -> bool:
    if name.startswith("."):
        return True
    if os.name == "nt":
        try:
            return bool(os.stat(full).st_file_attributes & 2)  # FILE_ATTRIBUTE_HIDDEN
        except (OSError, AttributeError):
            return False
    return False


def collect_files(source: str, opts: SearchOptions, exclude: Optional[str] = None) -> list:
    """列出要處理的檔案。exclude 通常是輸出檔,避免把自己的結果又掃進去。"""
    exts = {e.lower() for e in opts.exts}
    ex = os.path.normcase(os.path.abspath(exclude)) if exclude else None
    files = []
    for root, dirs, names in os.walk(source):
        if opts.skip_hidden:
            dirs[:] = [d for d in dirs if not _is_hidden(os.path.join(root, d), d)]
        dirs.sort()
        for n in sorted(names):
            if n.startswith("~$"):  # Excel 暫存檔
                continue
            if os.path.splitext(n)[1].lower() not in exts:
                continue
            full = os.path.join(root, n)
            if opts.skip_hidden and _is_hidden(full, n):
                continue
            if ex and os.path.normcase(os.path.abspath(full)) == ex:
                continue
            files.append(full)
        if not opts.recursive:
            break
    return files


# ───────────────────────────── 讀檔 ─────────────────────────────
def detect_csv_encoding(path: str) -> str:
    """只解碼不解析,先找出能完整讀完整個檔案的編碼。"""
    with open(path, "rb") as f:
        head = f.read(4)
    if head[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "utf-16"
    for enc in ("utf-8-sig", "cp950", "gb18030"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                while f.read(1 << 20):
                    pass
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def iter_frames(path: str) -> Iterator[tuple]:
    """逐塊產生 (工作表名稱或 None, DataFrame)。所有欄位都以字串讀入。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        enc = detect_csv_encoding(path)
        try:
            with pd.read_csv(path, dtype=str, encoding=enc, na_filter=False,
                             chunksize=CSV_CHUNK_ROWS) as reader:
                for chunk in reader:
                    yield None, chunk
        except pd.errors.EmptyDataError:
            return
    else:
        with pd.ExcelFile(path) as xl:
            for sheet in xl.sheet_names:
                yield sheet, pd.read_excel(xl, sheet_name=sheet, dtype=str, na_filter=False)


# ───────────────────────────── 比對 ─────────────────────────────
def _transform(s: pd.Series, opts: SearchOptions) -> pd.Series:
    if opts.normalize_width:
        s = s.str.normalize("NFKC")
    if not opts.case_sensitive:
        s = s.str.casefold()
    return s


def _transform_kw(kw: str, opts: SearchOptions) -> str:
    if opts.normalize_width:
        kw = unicodedata.normalize("NFKC", kw)
    if not opts.case_sensitive:
        kw = kw.casefold()
    return kw


def clean_keywords(opts: SearchOptions) -> list:
    kws = [str(k) for k in opts.keywords if str(k).strip()]
    if not kws:
        raise SearchError("請至少提供一個搜尋關鍵字。")
    return kws


def match_frame(df: pd.DataFrame, opts: SearchOptions):
    """回傳 (符合列的位置陣列, 每列命中的關鍵字文字)。"""
    keywords = clean_keywords(opts)
    clean = df.fillna("").astype(str)
    if clean.shape[0] == 0 or clean.shape[1] == 0:
        return np.array([], dtype=int), []

    masks = []
    if opts.whole and opts.whole_mode == "cell":
        # 任一儲存格的內容 (去頭尾空白後) 必須恰好等於關鍵字
        cells = clean.apply(lambda c: _transform(c.str.strip(), opts))
        for kw in keywords:
            k = _transform_kw(kw.strip(), opts)
            masks.append(cells.eq(k).to_numpy().any(axis=1))
    else:
        # 把整列串成一段文字再搜尋;分隔字元不是空白,所以關鍵字不會跨儲存格誤判
        text = clean.iloc[:, 0]
        for i in range(1, clean.shape[1]):
            text = text + _SEP + clean.iloc[:, i]
        text = _transform(text, opts)
        for kw in keywords:
            k = _transform_kw(kw, opts)
            if opts.whole:  # 完整單字:前後不可緊接文字
                pat = rf"(?<!\w){re.escape(k)}(?!\w)"
                masks.append(text.str.contains(pat, regex=True).to_numpy())
            else:
                masks.append(text.str.contains(k, regex=False).to_numpy())

    arr = np.column_stack(masks)
    hit = arr.all(axis=1) if opts.match_all else arr.any(axis=1)
    idx = np.flatnonzero(hit)
    labels = [", ".join(keywords[j] for j in np.flatnonzero(arr[i])) for i in idx]
    return idx, labels


def process_file(path: str, opts: SearchOptions) -> FileResult:
    """處理單一檔案 (會在子行程中執行,所以必須是模組層級函式)。"""
    try:
        name = os.path.basename(path)
        parts = []
        for sheet, df in iter_frames(path):
            if df.empty or df.shape[1] == 0:
                continue
            idx, labels = match_frame(df, opts)
            if len(idx) == 0:
                continue
            part = df.iloc[idx].copy()
            row_seq = df.index[idx].to_numpy() + 1   # 第幾筆資料 (不含標題列)
            part.columns = [f"{c}_原始" if str(c) in META_COLUMNS else str(c) for c in part.columns]
            part.reset_index(drop=True, inplace=True)
            meta = pd.DataFrame({
                "檔案路徑": path,
                "檔案名稱": name,
                "來源工作表": sheet or "",
                "資料列序": row_seq,
                "命中關鍵字": labels,
            })
            parts.append(pd.concat([meta, part], axis=1))
        if not parts:
            return FileResult(path)
        return FileResult(path, pd.concat(parts, ignore_index=True, sort=False))
    except Exception as e:  # noqa: BLE001
        return FileResult(path, error=str(e) or type(e).__name__)


# ───────────────────────────── 輸出 ─────────────────────────────
def with_timestamp(path: str) -> str:
    root, ext = os.path.splitext(path)
    return f"{root}_{time.strftime('%Y%m%d_%H%M%S')}{ext}"


def _unique(names: list) -> list:
    seen, out = {}, []
    for n in map(str, names):
        if n in seen:
            seen[n] += 1
            out.append(f"{n}.{seen[n]}")
        else:
            seen[n] = 0
            out.append(n)
    return out


def _display_width(s: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in s)


def write_output(final: pd.DataFrame, out: str, log: Callable = lambda *_: None) -> str:
    """寫出結果,回傳實際寫入的路徑 (超過 Excel 列數上限時會改存 CSV)。"""
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    final = final.copy()
    final.columns = _unique(list(final.columns))

    if out.lower().endswith(".xlsx") and len(final) > EXCEL_MAX_ROWS:
        out = os.path.splitext(out)[0] + ".csv"
        log("warn", f"結果超過 Excel 列數上限 ({EXCEL_MAX_ROWS:,} 列),改存為 CSV:{out}")

    if out.lower().endswith(".csv"):
        final.to_csv(out, index=False, encoding="utf-8-sig")
        return out

    for c in final.columns:  # openpyxl 不接受控制字元
        try:
            final[c] = final[c].str.replace(_ILLEGAL_XML, "", regex=True)
        except AttributeError:
            pass

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        final.to_excel(xw, index=False, sheet_name="結果")
        ws = xw.sheets["結果"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        head_font = Font(bold=True, color="FFFFFF")
        head_fill = PatternFill("solid", fgColor="4A66E8")
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 22
        for i, col in enumerate(final.columns, 1):
            sample = final[col].head(200).fillna("").astype(str)
            w = max([_display_width(str(col))] + [_display_width(v) for v in sample])
            ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 60)
    return out


# ───────────────────────────── 主流程 ─────────────────────────────
def _noop(*_a, **_k):
    return None


def run_search(
    source: str,
    output: str,
    opts: SearchOptions,
    *,
    workers: int = 1,
    on_progress: Callable = _noop,   # (已完成數, 總數, 目前檔案)
    on_stats: Callable = _noop,      # (已掃描, 命中檔案, 命中筆數)
    on_log: Callable = _noop,        # (等級, 訊息)
    should_cancel: Callable = lambda: False,
) -> SearchSummary:
    t0 = time.time()
    clean_keywords(opts)  # 提早驗證
    files = collect_files(source, opts, exclude=output)
    total = len(files)
    on_progress(0, total, "")

    def summary(**kw):
        base = dict(cancelled=False, scanned=0, matched_files=0, rows=0, elapsed=time.time() - t0)
        base.update(kw)
        return SearchSummary(**base)

    if total == 0:
        on_log("warn", "資料夾內沒有符合副檔名的檔案。")
        return summary()

    use_pool = workers > 1 and total >= POOL_MIN_FILES
    on_log("info", f"找到 {total} 個檔案,開始搜尋" + (f" (使用 {min(workers, total)} 個行程)…" if use_pool else "…"))

    results: dict = {}
    state = dict(scanned=0, files=0, rows=0)

    def handle(i: int, res: FileResult):
        results[i] = res
        state["scanned"] += 1
        if res.error:
            on_log("err", f"讀取失敗  {res.path}: {res.error}")
        elif res.frame is not None:
            state["files"] += 1
            state["rows"] += len(res.frame)
            on_log("ok", f"找到符合項目  {res.path}  ({len(res.frame)} 筆)")
        on_stats(state["scanned"], state["files"], state["rows"])
        on_progress(state["scanned"], total, res.path)

    def run_sequential(indices):
        for i in indices:
            if should_cancel():
                return
            on_progress(state["scanned"], total, files[i])
            handle(i, process_file(files[i], opts))

    if use_pool:
        try:
            _run_pool(files, opts, workers, should_cancel, handle)
        except BrokenProcessPool:
            on_log("warn", "多核心處理失敗,剩餘檔案改用單核心繼續。")
            run_sequential([i for i in range(total) if i not in results])
    else:
        run_sequential(range(total))

    on_progress(state["scanned"], total, "")
    if should_cancel():
        on_log("warn", "已取消,未寫入結果。")
        return summary(cancelled=True, scanned=state["scanned"],
                       matched_files=state["files"], rows=state["rows"])

    frames = [results[i].frame for i in sorted(results) if results[i].frame is not None]
    if not frames:
        on_log("warn", "未找到包含指定關鍵字的資料。")
        return summary(scanned=state["scanned"])

    final = pd.concat(frames, ignore_index=True, sort=False)
    try:
        out = write_output(final, output, on_log)
    except PermissionError:
        raise SearchError(f"無法寫入「{output}」。\n檔案可能正被 Excel 開啟,請先關閉後再試。") from None
    except OSError as e:
        raise SearchError(f"寫入結果失敗:{e}") from None
    on_log("ok", f"搜尋完成,結果已寫入:{out}")
    return summary(scanned=state["scanned"], matched_files=state["files"], rows=state["rows"],
                   output=out, preview=final.head(500).fillna(""))


def _run_pool(files, opts, workers, should_cancel, handle):
    # 一律用 spawn:跟 Windows 行為一致,也避免在有 Qt 執行緒的行程裡 fork
    ctx = multiprocessing.get_context("spawn")
    ex = ProcessPoolExecutor(max_workers=min(workers, len(files)), mp_context=ctx)
    futs = {ex.submit(process_file, p, opts): i for i, p in enumerate(files)}
    pending = set(futs)
    try:
        while pending:
            if should_cancel():
                break
            done, pending = wait(pending, timeout=0.2, return_when=FIRST_COMPLETED)
            for f in done:
                handle(futs[f], f.result())
    finally:
        ex.shutdown(wait=True, cancel_futures=True)


# ───────────────────────────── 命令列 ─────────────────────────────
def _cli():
    import argparse

    ap = argparse.ArgumentParser(description="搜尋資料夾內 Excel / CSV 中含關鍵字的整列資料")
    ap.add_argument("source", help="來源資料夾")
    ap.add_argument("output", help="輸出檔案 (.xlsx 或 .csv)")
    ap.add_argument("-k", "--keyword", action="append", required=True, help="關鍵字,可重複指定")
    ap.add_argument("--case", action="store_true", help="區分大小寫")
    ap.add_argument("--whole", choices=["cell", "word"], help="完整比對方式")
    ap.add_argument("--all", action="store_true", help="全部關鍵字都要出現")
    ap.add_argument("--width", action="store_true", help="全形半形視為相同")
    ap.add_argument("--no-recursive", action="store_true", help="不搜尋子資料夾")
    ap.add_argument("-j", "--jobs", type=int, default=0, help="行程數 (預設自動)")
    a = ap.parse_args()

    opts = SearchOptions(
        keywords=tuple(a.keyword), case_sensitive=a.case, whole=bool(a.whole),
        whole_mode=a.whole or "cell", match_all=a.all, normalize_width=a.width,
        recursive=not a.no_recursive)
    jobs = a.jobs or max(1, min(8, (os.cpu_count() or 2) - 1))
    try:
        r = run_search(a.source, a.output, opts, workers=jobs,
                       on_log=lambda lv, m: print(f"[{lv}] {m}"))
    except SearchError as e:
        raise SystemExit(str(e))
    print(f"掃描 {r.scanned} 個檔案,命中 {r.matched_files} 個檔案、{r.rows} 筆,耗時 {r.elapsed:.1f} 秒")


if __name__ == "__main__":
    multiprocessing.freeze_support()
    _cli()
